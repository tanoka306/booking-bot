"""
Телеграм-бот бронирования столов: Tanoka / Hit me baby
Хостинг: Bothost (Python), работает через polling — ответы мгновенные.
Без Google: все брони хранятся в Excel-файле, который бот держит
закреплённым в чате администратора и обновляет после каждой заявки.

Сценарий гостя:
  /start → кнопка заведения → схема зала (фото + текст) → имя →
  номер столика (кнопки) → дата и время → телефон →
  строка в таблице броней + уведомление администратору
"""

import asyncio
import io
import logging
import os
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from aiogram import Bot, Dispatcher, F, Router
from aiogram.filters import Command, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (
    BufferedInputFile,
    CallbackQuery,
    FSInputFile,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    InputMediaDocument,
    KeyboardButton,
    Message,
    ReplyKeyboardMarkup,
    ReplyKeyboardRemove,
)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

# ================== НАСТРОЙКИ ==================
# Задаются в панели Bothost → «Переменные окружения»
BOT_TOKEN = os.getenv("BOT_TOKEN", "")
ADMIN_CHAT_ID = os.getenv("ADMIN_CHAT_ID", "")  # чат администратора (узнать: /myid)
TIMEZONE = os.getenv("TIMEZONE", "Asia/Omsk")

BASE_DIR = Path(__file__).resolve().parent
IMAGES_DIR = BASE_DIR / "images"
FILE_NAME = "bookings.xlsx"

# Заведения: текст под схемой и номера столов (поменяйте под свои схемы)
VENUES = {
    "tanoka": {
        "name": "Tanoka",
        "plan_text": "Схема зала Tanoka",
        "tables": [str(i) for i in range(1, 11)],
    },
    "hit": {
        "name": "Hit me baby",
        "plan_text": "Схема зала Hit me baby",
        "tables": [str(i) for i in range(1, 11)],
    },
}

HEADERS = ["Создано", "Заведение", "Имя", "Номер столика",
           "Дата и время визита", "Телефон", "Telegram", "Chat ID"]
# ===============================================

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("bot")

router = Router()
photo_cache: dict[str, str] = {}  # file_id отправленных схем — повторно шлются мгновенно


class Booking(StatesGroup):
    name = State()
    table = State()
    datetime = State()
    phone = State()


# ---------- Таблица броней (Excel в Telegram) ----------

class BookingStore:
    """
    Брони живут в памяти и в файле bookings.xlsx, закреплённом в чате администратора.
    При перезапуске бот скачивает закреплённый файл и продолжает с того же места,
    поэтому данные не теряются даже на бесплатном тарифе Bothost.
    """

    def __init__(self):
        self.rows: list[list] = []
        self.message_id: int | None = None  # id закреплённого сообщения с файлом
        self.lock = asyncio.Lock()

    def to_bytes(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Брони"
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in self.rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        widths = [17, 14, 22, 14, 24, 18, 18, 14]
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def caption(self) -> str:
        now = datetime.now(ZoneInfo(TIMEZONE)).strftime("%d.%m.%Y %H:%M")
        return f"📋 Таблица броней — {len(self.rows)} шт.\nОбновлено: {now}\n\nНе открепляйте это сообщение."

    async def load(self, bot: Bot) -> None:
        """При старте забираем закреплённый файл из чата администратора."""
        if not ADMIN_CHAT_ID:
            log.warning("ADMIN_CHAT_ID не задан — брони не будут сохраняться между перезапусками")
            return
        try:
            chat = await bot.get_chat(ADMIN_CHAT_ID)
            pinned = chat.pinned_message
            if not (pinned and pinned.document and pinned.document.file_name == FILE_NAME):
                log.info("Закреплённой таблицы нет — начинаем новую")
                return
            buf = io.BytesIO()
            await bot.download(pinned.document.file_id, destination=buf)
            buf.seek(0)
            ws = load_workbook(buf).active
            self.rows = [
                ["" if v is None else str(v) for v in r]
                for r in ws.iter_rows(min_row=2, values_only=True)
                if any(v is not None for v in r)
            ]
            self.message_id = pinned.message_id
            log.info("Загружено броней из Telegram: %s", len(self.rows))
        except Exception as e:
            log.exception("Не удалось загрузить таблицу броней: %s", e)

    async def add(self, bot: Bot, row: list) -> bool:
        async with self.lock:
            self.rows.append(row)
            return await self._publish(bot)

    async def _publish(self, bot: Bot) -> bool:
        if not ADMIN_CHAT_ID:
            return False
        doc = BufferedInputFile(self.to_bytes(), filename=FILE_NAME)

        # 1) Тихо заменяем файл в закреплённом сообщении
        if self.message_id:
            try:
                await bot.edit_message_media(
                    chat_id=ADMIN_CHAT_ID,
                    message_id=self.message_id,
                    media=InputMediaDocument(media=doc, caption=self.caption()),
                )
                return True
            except Exception as e:
                log.warning("Не удалось обновить закреплённый файл, отправляю новый: %s", e)
                doc = BufferedInputFile(self.to_bytes(), filename=FILE_NAME)

        # 2) Нет закреплённого — отправляем и закрепляем
        try:
            msg = await bot.send_document(ADMIN_CHAT_ID, doc, caption=self.caption())
            self.message_id = msg.message_id
            try:
                await bot.pin_chat_message(ADMIN_CHAT_ID, msg.message_id, disable_notification=True)
            except Exception as e:
                log.warning("Не удалось закрепить файл (у бота нет прав?): %s", e)
            return True
        except Exception as e:
            log.exception("Не удалось отправить таблицу броней: %s", e)
            return False


store = BookingStore()


# ---------- Клавиатуры ----------

def venues_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="Забронировать стол в Tanoka", callback_data="venue:tanoka")],
        [InlineKeyboardButton(text="Забронировать стол в Hit me baby", callback_data="venue:hit")],
    ])


def tables_kb(venue_key: str) -> InlineKeyboardMarkup:
    tables = VENUES[venue_key]["tables"]
    rows = [
        [InlineKeyboardButton(text=f"№{t}", callback_data=f"table:{t}") for t in tables[i:i + 5]]
        for i in range(0, len(tables), 5)
    ]
    return InlineKeyboardMarkup(inline_keyboard=rows)


def phone_kb() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        keyboard=[[KeyboardButton(text="Отправить мой номер", request_contact=True)]],
        resize_keyboard=True,
        one_time_keyboard=True,
    )


# ---------- Вспомогательное ----------

def is_admin_chat(message: Message) -> bool:
    return bool(ADMIN_CHAT_ID) and str(message.chat.id) == str(ADMIN_CHAT_ID)


def find_plan_image(venue_key: str) -> Path | None:
    # Ищем и в папке images, и рядом с main.py, имя файла в любом регистре
    for folder in (IMAGES_DIR, BASE_DIR):
        if not folder.is_dir():
            continue
        for p in folder.iterdir():
            if p.is_file() and p.stem.lower() == venue_key and \
                    p.suffix.lower() in (".jpg", ".jpeg", ".png", ".webp"):
                return p
    return None


async def notify_admin(bot: Bot, text: str) -> None:
    if not ADMIN_CHAT_ID:
        return
    try:
        await bot.send_message(ADMIN_CHAT_ID, text)
    except Exception as e:
        log.warning("Не удалось написать администратору: %s", e)


async def send_plan(message: Message, venue_key: str) -> None:
    """Отправляет картинку схемы зала с текстом."""
    venue = VENUES[venue_key]
    caption = venue["plan_text"]

    file_id = photo_cache.get(venue_key)
    if file_id:
        try:
            await message.answer_photo(file_id, caption=caption)
            return
        except Exception:
            photo_cache.pop(venue_key, None)

    path = find_plan_image(venue_key)
    if path:
        try:
            sent = await message.answer_photo(FSInputFile(path), caption=caption)
            photo_cache[venue_key] = sent.photo[-1].file_id
            return
        except Exception as e:
            log.exception("Не удалось отправить схему %s: %s", venue_key, e)

    await message.answer(caption)
    await notify_admin(message.bot, f"⚠️ Схема зала {venue['name']} не найдена: "
                                    f"положите файл images/{venue_key}.jpg в проект")


# ---------- Команды ----------

@router.message(CommandStart())
async def cmd_start(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("Здравствуйте! Где хотите забронировать стол?",
                         reply_markup=ReplyKeyboardRemove())
    await message.answer("Выберите заведение:", reply_markup=venues_kb())


@router.message(Command("myid"))
async def cmd_myid(message: Message):
    await message.answer(f"Chat ID: {message.chat.id}")


# Только для администратора: прислать актуальную таблицу
@router.message(Command("export"))
async def cmd_export(message: Message):
    if not is_admin_chat(message):
        return
    await message.answer_document(
        BufferedInputFile(store.to_bytes(), filename=FILE_NAME),
        caption=f"Всего броней: {len(store.rows)}",
    )


# ---------- Сценарий бронирования ----------

# Шаг 1: выбор заведения → схема зала → вопрос имени
@router.callback_query(F.data.startswith("venue:"))
async def choose_venue(callback: CallbackQuery, state: FSMContext):
    venue_key = callback.data.split(":", 1)[1]
    await callback.answer()
    if venue_key not in VENUES:
        return
    await state.clear()
    await state.update_data(venue=venue_key)
    await send_plan(callback.message, venue_key)
    await callback.message.answer("Ваше имя")
    await state.set_state(Booking.name)


# Шаг 2: имя → выбор столика
@router.message(Booking.name, F.text)
async def get_name(message: Message, state: FSMContext):
    await state.update_data(name=message.text.strip())
    data = await state.get_data()
    await message.answer("Выберите номер столика (номера указаны на схеме зала)",
                         reply_markup=tables_kb(data["venue"]))
    await state.set_state(Booking.table)


# Шаг 3: столик (кнопкой или текстом) → дата и время
async def _set_table(message: Message, state: FSMContext, table: str):
    await state.update_data(table=table)
    await message.answer(f"Столик №{table} выбран.\n\nДата и время визита")
    await state.set_state(Booking.datetime)


@router.callback_query(Booking.table, F.data.startswith("table:"))
async def choose_table_button(callback: CallbackQuery, state: FSMContext):
    await callback.answer()
    await _set_table(callback.message, state, callback.data.split(":", 1)[1])


@router.message(Booking.table, F.text)
async def choose_table_text(message: Message, state: FSMContext):
    await _set_table(message, state, message.text.strip().lstrip("№").strip())


# Шаг 4: дата и время → телефон
@router.message(Booking.datetime, F.text)
async def get_datetime(message: Message, state: FSMContext):
    await state.update_data(datetime=message.text.strip())
    await message.answer("Контактный телефон", reply_markup=phone_kb())
    await state.set_state(Booking.phone)


# Шаг 5: телефон → сохранение
@router.message(Booking.phone, F.contact | F.text)
async def get_phone(message: Message, state: FSMContext):
    phone = message.contact.phone_number if message.contact else message.text.strip()
    data = await state.get_data()
    await state.clear()

    venue = VENUES[data["venue"]]["name"]
    user = message.from_user
    username = f"@{user.username}" if user and user.username else ""
    created = datetime.now(ZoneInfo(TIMEZONE)).strftime("%d.%m.%Y %H:%M")

    await message.answer(
        f"Спасибо! Заявка на бронь в {venue} принята.\n"
        "Мы свяжемся с вами для подтверждения.\n\nНовая бронь: /start",
        reply_markup=ReplyKeyboardRemove(),
    )

    await notify_admin(message.bot,
        "🔔 Новая бронь\n\n"
        f"Заведение: {venue}\n"
        f"Имя: {data['name']}\n"
        f"Столик: №{data['table']}\n"
        f"Дата и время: {data['datetime']}\n"
        f"Телефон: {phone}\n"
        + (f"Telegram: {username}\n" if username else "")
    )

    saved = await store.add(message.bot, [
        created, venue, data["name"], data["table"], data["datetime"],
        phone, username, str(message.chat.id),
    ])
    if not saved:
        await notify_admin(message.bot, "⚠️ Не удалось обновить таблицу броней — смотрите логи Bothost")


# Если гость прислал не то (стикер, фото и т.п.)
@router.message(Booking.name)
@router.message(Booking.datetime)
async def need_text(message: Message):
    await message.answer("Пожалуйста, ответьте текстом.")


@router.message(Booking.phone)
async def need_phone(message: Message):
    await message.answer("Напишите номер или нажмите кнопку «Отправить мой номер».",
                         reply_markup=phone_kb())


# Любое сообщение вне сценария → предлагаем начать (кроме чата администратора)
@router.message()
async def fallback(message: Message):
    if is_admin_chat(message) and message.chat.type != "private":
        return
    await message.answer("Выберите заведение:", reply_markup=venues_kb())


# ---------- Запуск ----------

async def main():
    if not BOT_TOKEN:
        raise SystemExit("Не задан BOT_TOKEN в переменных окружения")

    bot = Bot(BOT_TOKEN)
    dp = Dispatcher(storage=MemoryStorage())
    dp.include_router(router)

    await store.load(bot)

    for key in VENUES:
        if not find_plan_image(key):
            log.warning("Нет картинки схемы: images/%s.jpg", key)

    await bot.delete_webhook(drop_pending_updates=True)  # отключаем старый webhook, если был
    log.info("Бот запущен")
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
